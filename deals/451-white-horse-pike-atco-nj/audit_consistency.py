"""Independent recomputation + cross-document consistency audit."""
from openpyxl import load_workbook
from pypdf import PdfReader
import re

def irr(cf):
    lo,hi=-0.95,2.0
    f=lambda r: sum(c/(1+r)**i for i,c in enumerate(cf))
    for _ in range(400):
        m=(lo+hi)/2
        if f(lo)*f(m)<=0: hi=m
        else: lo=m
    return (lo+hi)/2

# ---------- 1. INDEPENDENT MODEL ----------
P, ASK, SF, RENT, OPT = 1_100_000, 1_600_000, 4400, 132_000, 154_000
TAX, INS, CAM = 39_402, 12_000, 10_000
MKT = 22
S = [("renew",0.40,OPT,0.0925,25_000),("cannabis",0.10,132_000,0.0950,325_000),
     ("bank",0.20,96_800,0.0875,275_000),("medical",0.20,105_600,0.0825,500_000),
     ("qsr",0.05,92_400,0.0700,330_000),("dark",0.05,0,0,123_000)]
net={}
for n,p,noi,cap,cost in S:
    stab = 750_000 if cap==0 else noi/cap
    net[n]=stab-cost
W = sum(p*net[n] for n,p,_,_,_ in S)
cf  = lambda px,rev: [-px]+[RENT]*6+[RENT*0.5+rev]
EXP = {"ours":irr(cf(P,W)), "ask":irr(cf(ASK,W))}
IND = {n:(irr(cf(P,net[n])), irr(cf(ASK,net[n]))) for n,_,_,_,_ in S}
VALS = {"in_place":RENT/0.0975, "assessor":760_200/0.5981,
        "psf":273*SF, "mkt_rent":MKT*SF/0.0875}
WTS  = {"in_place":0.40,"assessor":0.20,"psf":0.15,"mkt_rent":0.25}
VALIND = sum(VALS[k]*WTS[k] for k in VALS)/sum(WTS.values())
LAND_NOI, LAND_COST, LAND_CAP = 55_000+20_000, 150_000, 0.0925
LAND_VC = LAND_NOI/LAND_CAP - LAND_COST

EXPECT = {
 "price_psf": P/SF, "cap_ours": RENT/P, "cap_ask": RENT/ASK,
 "irr_exp_ours": EXP["ours"], "irr_exp_ask": EXP["ask"],
 "irr_renew_ours": IND["renew"][0], "irr_dark_ours": IND["dark"][0],
 "irr_dark_ask": IND["dark"][1],
 "wrev": W, "valind": VALIND, "mkt_rent_val": VALS["mkt_rent"],
 "payback_pct": RENT*6.5/P, "equity_mult": (RENT*6.5+W)/P,
 "land_noi": LAND_NOI, "land_vc": LAND_VC,
 "yoc_land": (RENT+LAND_NOI)/P, "dark_plus_land": (18*SF+LAND_NOI)/P,
 "breakeven_psf": P*0.085/SF, "excess_ac": 3.13-SF/43560-0.55,
}
print("="*78); print("1. INDEPENDENT RECOMPUTATION"); print("="*78)
for k,v in EXPECT.items(): print(f"  {k:20s} {v:,.4f}")

# ---------- 2. WORKBOOK ----------
wb = load_workbook("451_White_Horse_Pike_Acquisition_Model.xlsx", data_only=True)
def cell(sheet,label,col=2,lim=95):
    ws=wb[sheet]
    for r in range(1,lim):
        if str(ws.cell(r,1).value or "").strip()==label: return ws.cell(r,col).value
def cell2(sheet,label,col,lim=95):
    ws=wb[sheet]
    for r in range(1,lim):
        if str(ws.cell(r,2).value or "").strip().startswith(label): return ws.cell(r,col).value
GOT = {
 "price_psf": cell("Returns","Price per square foot"),
 "cap_ours": cell("Returns","Going-in cap rate"),
 "cap_ask": cell("Executive Summary","Going-in cap rate",3),
 "irr_exp_ours": cell("Returns","Unlevered IRR (expected)"),
 "irr_exp_ask": cell("Executive Summary","Expected unlevered IRR",3),
 "irr_renew_ours": cell("Executive Summary","Best case IRR (tenant renews)"),
 "irr_dark_ours": cell("Executive Summary","Worst case IRR (total vacancy)"),
 "irr_dark_ask": cell("Executive Summary","Worst case IRR (total vacancy)",3),
 "wrev": cell("Reversion Scenarios","PROBABILITY-WEIGHTED REVERSION",7),
 "valind": cell("Valuation","WEIGHTED VALUE INDICATION",3),
 "mkt_rent_val": cell("Valuation","Absolute floor (market-rent value)"),
 "payback_pct": cell("Returns","  ... as % of purchase price"),
 "equity_mult": cell("Returns","Equity multiple"),
 "land_noi": cell2("Land Upside","COMBINED INCREMENTAL NOI",3),
 "land_vc": cell2("Land Upside","Value created (net of build-out",3),
 "yoc_land": cell("Land Upside","Yield on cost"),
 "dark_plus_land": cell("Land Upside","DOWNSIDE RE-TEST: dark building @ $18/SF + land income"),
 "breakeven_psf": cell("Valuation","Break-even rent to hold our basis"),
 "excess_ac": cell("Land Upside","DEVELOPABLE EXCESS LAND"),
}
print("\n"+"="*78); print("2. WORKBOOK vs INDEPENDENT MODEL"); print("="*78)
bad=0
for k,exp in EXPECT.items():
    got=GOT.get(k)
    if got is None: print(f"  [MISSING] {k}"); bad+=1; continue
    ok = abs(float(got)-exp) < max(abs(exp)*0.005, 0.0001)
    if not ok: bad+=1
    print(f"  {'OK ' if ok else 'MISMATCH'} {k:20s} model={exp:>14,.4f}  wb={float(got):>14,.4f}")
print(f"\n  -> {bad} mismatch(es)")

# ---------- 3. PDF ----------
txt=" ".join((p.extract_text() or "") for p in PdfReader("451_White_Horse_Pike_Analysis.pdf").pages)
txt=re.sub(r"\s+"," ",txt)
print("\n"+"="*78); print("3. PDF FIGURE CHECK"); print("="*78)
probes=[("$1,100,000",1),("$250.00",1),("12.00%",1),("4.33%",1),("15.80%",1),("+6.20%",1),
        ("−1.43%",0),("$1,252,500",1),("$1,106,3",1),("$1,165,500",1),("$858,000",1),
        ("1.84x",1),("$75,000",1),("$661,000",1),("18.8%",1),("14.0%",1),("$21.25",1),
        ("2.48",1),("$22.00",1),("$27.43",1),("9.35%",1),("8.87%",1)]
miss=[]
for s_,_ in probes:
    hit = s_ in txt
    if not hit: miss.append(s_)
    print(f"  {'OK ' if hit else 'ABSENT'} {s_}")
print(f"\n  -> absent: {miss}")
