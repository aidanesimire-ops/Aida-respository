"""
tab_income.py — ASSEMBLAGE INCOME VALUATION (consolidated cash flows).
Institutional build: consolidate every asset's cash flow, separate as-is vs.
stabilized NOI, value by direct-cap + DCF, size debt to the LESSER OF LTV /
DSCR / debt-yield, run returns at BOTH the income price and the covered-land
price, quantify the land-appreciation break-even, and stress it with two
sensitivity grids. This tab is the single source of truth for consolidated
(portfolio) returns — computed from actual cash flows, never averaged IRRs.
"""
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule
from mblib import (F_ACCT, F_ACCT_TOP, F_PCT1, F_PCT2, F_MULT, F_PSF, F_NUM, F_YR)

L = 2
ACQ = 3
def pc(y): return 3 + y
def CL(c): return get_column_letter(c)

ASSETS = ["Shahidi Retail", "Publix & Starbucks", "Sunrise Plaza", "Office Condo", "Land"]
LABEL = {"Shahidi Retail": "Shahidi Retail (Galleria Plaza)",
         "Publix & Starbucks": "Publix + Starbucks",
         "Sunrise Plaza": "Sunrise Plaza (Kar Luen)",
         "Office Condo": "Galleria Corporate Centre (office)",
         "Land": "1040 Bayview (interim office)"}
HAS_CAP = {"Shahidi Retail": True, "Publix & Starbucks": True, "Sunrise Plaza": True,
           "Office Condo": True, "Land": False}
ASIS = {"Shahidi Retail": "INPLACE_NOI", "Publix & Starbucks": "AS_IS_NOI",
        "Sunrise Plaza": "AS_IS_NOI", "Office Condo": "AS_IS_NOI", "Land": "INOI"}
STAB = {"Shahidi Retail": "STABNOI", "Publix & Starbucks": "STABNOI",
        "Sunrise Plaza": "STABNOI", "Office Condo": "STABNOI", "Land": "INOI"}
HEAT = ColorScaleRule(start_type="min", start_color="F8696B",
                      mid_type="percentile", mid_value=50, mid_color="FFEB84",
                      end_type="max", end_color="63BE7B")


def _row(coord):
    return int("".join(c for c in coord if c.isdigit()))


def build(s, regs):
    def link(sheet, anchor, y):
        return f"'{sheet}'!{CL(pc(y))}{_row(regs[sheet][anchor])}"
    def cell(sheet, name):
        return f"'{sheet}'!{regs[sheet][name]}"
    s.colw({"A": 2.5, "B": 36, "C": 14, "D": 13, "E": 14, "F": 13, "G": 13,
            "H": 13, "I": 13, "J": 13, "K": 13, "L": 13, "M": 13})
    r = 1
    s.put(r, L, "ASSEMBLAGE — INCOME VALUATION  (CONSOLIDATED CASH FLOWS)", style="banner", align="left", merge=(r, 13)); s.rowh(r, 26); r += 1
    s.put(r, L, "Combine every asset's cash flow → value the block on the income it produces (as-is + stabilized + DCF) → "
                "contrast with the covered-land / highest-and-best-use price. Portfolio returns run off consolidated cash flows.",
          style="banner_sub", align="left", merge=(r, 13)); s.rowh(r, 30); r += 2

    s.section(r, L, 13, "HOW THIS DIFFERS FROM THE HBU ASSEMBLAGE"); r += 1
    for t in [
        "The Assemblage tab prices the block on a highest-and-best-use / covered-land basis (land value + assemblage premium to control every",
        "parcel). This tab prices the SAME block on the income it produces. Income-producing real estate is worth the rent it throws off — which",
        "is less than the land-control price. Returns are computed by consolidating the actual cash flows and running one IRR (not by averaging",
        "asset IRRs, which is mathematically invalid). Debt is sized to the lesser of LTV / DSCR / debt-yield, the way a lender underwrites.",
    ]:
        s.put(r, L, t, style="calc", align="left", merge=(r, 13)); r += 1
    r += 1

    # ================= consolidated cash flow =================
    s.section(r, L, 13, "CONSOLIDATED CASH FLOW  —  10-year annual  (🟢 links to each asset tab)"); r += 1
    s.put(r, L, "Year", style="subhead", align="left")
    for y in range(1, 11): s.put(r, pc(y), y, style="subhead", align="center", fmt=F_YR)
    r += 1
    s.put(r, L, "Year ending", style="note", align="left")
    for y in range(1, 11): s.put(r, pc(y), f"Dec-{2026+y}", style="note", align="center")
    r += 1
    s.put(r, L, "NET OPERATING INCOME by asset", style="subhead", align="left", merge=(r, 13)); r += 1
    noi_first = r
    for a in ASSETS:
        s.put(r, L, f"  {LABEL[a]}", style="label", align="left")
        for y in range(1, 11):
            s.put(r, pc(y), f"={link(a, 'ROW_NOI', y)}", style="calc", color="008000",
                  fmt=(F_ACCT_TOP if y == 1 else F_ACCT), align="right")
        r += 1
    noi_last = r - 1
    noi_tot = r
    s.put(r, L, "Consolidated NOI", style="subtotal", align="left")
    for y in range(1, 11):
        s.put(r, pc(y), f"=SUM({CL(pc(y))}{noi_first}:{CL(pc(y))}{noi_last})",
              style="calc", fmt=(F_ACCT_TOP if y == 1 else F_ACCT), align="right", bold=True)
    r += 1
    cap_row = r
    s.put(r, L, "Less: consolidated capital (reserves + leasing)", style="label", align="left")
    for y in range(1, 11):
        terms = "+".join(link(a, "ROW_CAP", y) for a in ASSETS if HAS_CAP[a])
        s.put(r, pc(y), f"=-({terms})", style="calc", color="008000", fmt=F_ACCT, align="right")
    r += 1
    op_row = r
    s.put(r, L, "Consolidated unlevered operating CF", style="subtotal", align="left")
    for y in range(1, 11):
        s.put(r, pc(y), f"={CL(pc(y))}{noi_tot}+{CL(pc(y))}{cap_row}",
              style="calc", fmt=(F_ACCT_TOP if y == 1 else F_ACCT), align="right", bold=True)
    r += 1
    rev_row = r
    s.put(r, L, "Consolidated reversion (net sale, exit yr)", style="label", align="left")
    for y in range(1, 11):
        s.put(r, pc(y), f"={'+'.join(link(a, 'ROW_REV', y) for a in ASSETS)}",
              style="calc", color="008000", fmt=F_ACCT, align="right")
    r += 2

    g = s.reg
    def R(n): return g[n]

    # ================= as-is / stabilized NOI =================
    s.section(r, L, 13, "AS-IS  vs.  STABILIZED NOI  (going-in vs. lease-up complete)"); r += 1
    s.put(r, L, "Consolidated AS-IS in-place NOI", style="subtotal", align="left")
    s.put(r, 3, f"={'+'.join(cell(a, ASIS[a]) for a in ASSETS)}", style="calc", color="008000",
          fmt=F_ACCT_TOP, align="right", name="ASIS_NOI", bold=True)
    s.put(r, 4, "current occupancy, no lease-up (going-in)", style="note", align="left", merge=(r, 13)); r += 1
    s.put(r, L, "Consolidated STABILIZED NOI (lease-up complete)", style="subtotal", align="left")
    s.put(r, 3, f"={'+'.join(cell(a, STAB[a]) for a in ASSETS)}", style="calc", color="008000",
          fmt=F_ACCT_TOP, align="right", name="STAB_NOI", bold=True)
    s.put(r, 4, "post lease-up (Shahidi to 95%, Sunrise to 95%, office to 88%)", style="note", align="left", merge=(r, 13)); r += 1
    s.put(r, L, "Value creation (lease-up NOI gain)", style="label", align="left")
    s.put(r, 3, f"={R('STAB_NOI')}-{R('ASIS_NOI')}", style="calc", fmt=F_ACCT, align="right"); r += 2

    # ================= valuation assumptions =================
    s.section(r, L, 13, "INCOME-VALUATION ASSUMPTIONS  —  🔵 blue inputs"); r += 1
    def inp(name, label, val, fmt, note):
        nonlocal r
        s.put(r, L, label, style="label", align="left")
        s.put(r, 3, val, style="input", fmt=fmt, align="right", name=name)
        s.put(r, 4, note, style="note", align="left", merge=(r, 13)); r += 1
    inp("CAP", "Blended going-in cap (as-is)", 0.0650, F_PCT2, "🔶 NOI-weighted retail/grocery/office")
    inp("SCAP", "Blended stabilized / exit cap", 0.0675, F_PCT2, "🔶 going-in + 25 bps")
    inp("TYLD", "Target unlevered yield (DCF discount)", 0.085, F_PCT1, "🔶 unlevered hurdle")
    inp("GROW", "Blended NOI growth (for reversion/sens.)", 0.025, F_PCT1, "🔶")
    inp("HOLD", "Hold period (yrs)", 5, F_YR, "🔶 assets exit together")
    inp("COS", "Cost of sale at exit", 0.02, F_PCT1, "🔶")
    inp("DOCSTAMP", "Doc-stamp / transfer tax (% price)", 0.0070, F_PCT2, "🔶 FL Broward $0.70/$100")
    inp("TITLE", "Title insurance (% price)", 0.0050, F_PCT2, "🔶")
    inp("LEGALDD", "Legal & due diligence (% price)", 0.0040, F_PCT2, "🔶")
    inp("ORIG", "Loan origination fee (% loan)", 0.0100, F_PCT2, "🔶")
    inp("LTV", "Max senior LTV", 0.60, F_PCT1, "🔶 constraint 1")
    inp("DSCRMIN", "Min DSCR (sizing constraint)", 1.30, F_MULT, "🔶 constraint 2")
    inp("DYMIN", "Min debt yield (sizing constraint)", 0.085, F_PCT1, "🔶 constraint 3")
    inp("RATE", "Senior rate", 0.065, F_PCT2, "🔶")
    inp("AMORT", "Amortization (yrs)", 30, F_YR, "🔶")
    r += 1

    # helper: mortgage constant, sums
    s.section(r, L, 13, "SIZING HELPERS"); r += 1
    def der(name, label, formula, fmt, note=""):
        nonlocal r
        s.put(r, L, label, style="label", align="left")
        s.put(r, 3, formula, style="calc", fmt=fmt, align="right", name=name)
        if note: s.put(r, 4, note, style="note", align="left", merge=(r, 13))
        r += 1
    der("MRATE", "Monthly rate", f"={R('RATE')}/12", F_PCT2)
    der("MCONST", "Mortgage constant (annual)",
        f"=({R('MRATE')}/(1-(1+{R('MRATE')})^(-{R('AMORT')}*12)))*12", F_PCT2, "DS ÷ loan")
    der("CLOSE", "Total transaction cost (% price)",
        f"={R('DOCSTAMP')}+{R('TITLE')}+{R('LEGALDD')}", F_PCT2, "doc stamps + title + legal/DD")
    der("SUMOP5", "Σ unlevered operating CF, Yr 1–hold",
        f"=SUM({CL(pc(1))}{op_row}:INDEX({CL(pc(1))}{op_row}:{CL(pc(10))}{op_row},{R('HOLD')}))", F_ACCT_TOP)
    der("EXITNOI", "Exit-year consolidated NOI",
        f"=INDEX({CL(pc(1))}{noi_tot}:{CL(pc(10))}{noi_tot},{R('HOLD')})", F_ACCT_TOP)
    der("REV_BASE", "Base-case consolidated reversion",
        f"=INDEX({CL(pc(1))}{rev_row}:{CL(pc(10))}{rev_row},{R('HOLD')})", F_ACCT_TOP, "sum of asset reversions")
    r += 1

    # ================= income valuation =================
    s.section(r, L, 13, "PURCHASE PRICE ON THE BASIS OF ASSEMBLAGE CASH FLOWS"); r += 1
    der("VAL_ASIS", "① As-is value  (as-is NOI ÷ going-in cap)", f"={R('ASIS_NOI')}/{R('CAP')}", F_ACCT_TOP, "going-in income value")
    der("VAL_STAB", "② Stabilized value  (stab NOI ÷ exit cap)", f"={R('STAB_NOI')}/{R('SCAP')}", F_ACCT_TOP, "post lease-up")
    # DCF valuation CF
    s.put(r, L, "Valuation CF (op CF ≤ hold + reversion)", style="label", align="left")
    for y in range(1, 11):
        s.put(r, pc(y), f"=IF({y}<={R('HOLD')},{CL(pc(y))}{op_row},0)+IF({y}={R('HOLD')},{CL(pc(y))}{rev_row},0)",
              style="calc", fmt=(F_ACCT_TOP if y == 1 else F_ACCT), align="right")
    valcf_row = r; r += 1
    der("VAL_DCF", "③ DCF value  (PV @ target unlevered yield)",
        f"=NPV({R('TYLD')},{CL(pc(1))}{valcf_row}:{CL(pc(10))}{valcf_row})", F_ACCT_TOP, "price where unlevered IRR = target")
    der("PRICE", "CONCLUDED INCOME-BASED PURCHASE PRICE", f"={R('VAL_ASIS')}", F_ACCT_TOP, "🔶 as-is direct cap; ②/③ cross-check")
    s.ws[f"B{r-1}"].font = s.ws[f"B{r-1}"].font.copy(bold=True)
    r += 1

    # ================= debt sizing (lesser-of) =================
    s.section(r, L, 13, "SENIOR DEBT SIZING  —  lesser of LTV / DSCR / debt-yield  (sized on as-is NOI)"); r += 1
    der("L_LTV", "Loan by max LTV", f"={R('PRICE')}*{R('LTV')}", F_ACCT_TOP)
    der("L_DSCR", "Loan by min DSCR", f"={R('ASIS_NOI')}/({R('DSCRMIN')}*{R('MCONST')})", F_ACCT_TOP)
    der("L_DY", "Loan by min debt yield", f"={R('ASIS_NOI')}/{R('DYMIN')}", F_ACCT_TOP)
    der("LOAN", "SIZED SENIOR LOAN (binding constraint)", f"=MIN({R('L_LTV')},{R('L_DSCR')},{R('L_DY')})", F_ACCT_TOP)
    s.put(r-1, 4, f'=IF({R("LOAN")}={R("L_LTV")},"LTV-constrained",IF({R("LOAN")}={R("L_DSCR")},"DSCR-constrained","Debt-yield-constrained"))',
          style="note", align="left", merge=(r-1, 13))
    der("DS", "Annual debt service", f"={R('LOAN')}*{R('MCONST')}", F_ACCT_TOP)
    der("LBAL5", "Loan balance at exit",
        f"={R('LOAN')}*(1+{R('MRATE')})^(12*{R('HOLD')})-({R('DS')}/12)*((1+{R('MRATE')})^(12*{R('HOLD')})-1)/{R('MRATE')}", F_ACCT_TOP)
    der("EQ", "Equity required (at income price)",
        f"={R('PRICE')}*(1+{R('CLOSE')})+{R('ORIG')}*{R('LOAN')}-{R('LOAN')}", F_ACCT_TOP)
    r += 1

    # ================= sources & uses =================
    s.section(r, L, 13, "SOURCES  &  USES  (at the income-based price)"); r += 1
    s.put(r, L, "USES", style="subhead", align="left", merge=(r, 3))
    s.put(r, 5, "SOURCES", style="subhead", align="left", merge=(r, 6)); r += 1
    u0 = r
    s.put(r, L, "  Purchase price", style="label", align="left")
    s.put(r, 3, f"={R('PRICE')}", style="calc", fmt=F_ACCT_TOP, align="right")
    s.put(r, 5, "  Senior loan (sized)", style="label", align="left")
    s.put(r, 6, f"={R('LOAN')}", style="calc", fmt=F_ACCT_TOP, align="right", name="SU_DEBT"); r += 1
    s.put(r, L, "  Doc-stamp / transfer tax", style="label", align="left")
    s.put(r, 3, f"={R('PRICE')}*{R('DOCSTAMP')}", style="calc", fmt=F_ACCT, align="right")
    s.put(r, 5, "  Sponsor equity (plug)", style="label", align="left")
    s.put(r, 6, f"={R('EQ')}", style="calc", fmt=F_ACCT_TOP, align="right", name="SU_EQ"); r += 1
    s.put(r, L, "  Title insurance", style="label", align="left")
    s.put(r, 3, f"={R('PRICE')}*{R('TITLE')}", style="calc", fmt=F_ACCT, align="right"); r += 1
    s.put(r, L, "  Legal & due diligence", style="label", align="left")
    s.put(r, 3, f"={R('PRICE')}*{R('LEGALDD')}", style="calc", fmt=F_ACCT, align="right"); r += 1
    s.put(r, L, "  Loan origination fee", style="label", align="left")
    s.put(r, 3, f"={R('ORIG')}*{R('LOAN')}", style="calc", fmt=F_ACCT, align="right"); r += 1
    uN = r - 1
    s.put(r, L, "TOTAL USES", style="subtotal", align="left")
    s.put(r, 3, f"=SUM({CL(3)}{u0}:{CL(3)}{uN})", style="calc", fmt=F_ACCT_TOP, align="right", bold=True, name="SU_USES")
    s.put(r, 5, "TOTAL SOURCES", style="subtotal", align="left")
    s.put(r, 6, f"={R('SU_DEBT')}+{R('SU_EQ')}", style="calc", fmt=F_ACCT_TOP, align="right", bold=True, name="SU_SRC"); r += 1
    s.put(r, L, "Check (sources − uses)", style="note", align="left")
    s.put(r, 3, f"={R('SU_SRC')}-{R('SU_USES')}", style="calc", fmt=F_ACCT, align="right")
    s.put(r, 5, "Loan-to-cost (LTC)", style="note", align="left")
    s.put(r, 6, f"={R('SU_DEBT')}/{R('SU_USES')}", style="calc", fmt=F_PCT1, align="right"); r += 2

    # ================= dual-scenario returns =================
    s.section(r, L, 13, "PORTFOLIO RETURNS  —  consolidated cash-flow IRR at income price vs. covered-land price"); r += 1
    px_hbu = cell("Assemblage", "ACQ")
    # per-scenario sized loan (DSCR/DY loans are price-independent; only LTV loan scales)
    s.put(r, L, "Scenario", style="subhead", align="left")
    s.put(r, 3, "① Income price", style="subhead", align="center")
    s.put(r, 5, "② Covered-land (HBU)", style="subhead", align="center", merge=(r, 6)); r += 1
    def two(name_i, name_h, label, f_income, f_hbu, fmt):
        nonlocal r
        s.put(r, L, label, style="label", align="left")
        s.put(r, 3, f_income, style="calc", fmt=fmt, align="right", name=name_i)
        s.put(r, 5, f_hbu, style="calc", fmt=fmt, align="right", name=name_h, merge=(r, 6))
        r += 1
    two("PX_I", "PX_H", "Purchase price", f"={R('PRICE')}", f"={px_hbu}", F_ACCT_TOP)
    two("LN_I", "LN_H", "Sized senior loan",
        f"={R('LOAN')}", f"=MIN({px_hbu}*{R('LTV')},{R('L_DSCR')},{R('L_DY')})", F_ACCT_TOP)
    two("EQ_I", "EQ_H", "Equity required",
        f"={R('PRICE')}*(1+{R('CLOSE')})+{R('ORIG')}*{R('LN_I')}-{R('LN_I')}",
        f"={px_hbu}*(1+{R('CLOSE')})+{R('ORIG')}*{R('LN_H')}-{R('LN_H')}", F_ACCT_TOP)
    two("GI_I", "GI_H", "Going-in cap (as-is NOI ÷ price)",
        f"={R('ASIS_NOI')}/{R('PRICE')}", f"={R('ASIS_NOI')}/{px_hbu}", F_PCT2)
    # unlevered CF rows (differ only in acq)
    ui_row = r
    s.put(r, L, "  Unlevered CF ① (acq + valuation CF)", style="note", align="left")
    s.put(r, ACQ, f"=-{R('PRICE')}*(1+{R('CLOSE')})", style="calc", fmt=F_ACCT, align="right")
    for y in range(1, 11): s.put(r, pc(y), f"={CL(pc(y))}{valcf_row}", style="calc", fmt=F_ACCT, align="right")
    r += 1
    uh_row = r
    s.put(r, L, "  Unlevered CF ② (acq + valuation CF)", style="note", align="left")
    s.put(r, ACQ, f"=-{px_hbu}*(1+{R('CLOSE')})", style="calc", fmt=F_ACCT, align="right")
    for y in range(1, 11): s.put(r, pc(y), f"={CL(pc(y))}{valcf_row}", style="calc", fmt=F_ACCT, align="right")
    r += 1
    # levered CF rows
    li_row = r
    s.put(r, L, "  Levered CF ① (equity)", style="note", align="left")
    s.put(r, ACQ, f"=-{R('EQ_I')}", style="calc", fmt=F_ACCT, align="right")
    for y in range(1, 11):
        s.put(r, pc(y), f"=IF({y}<={R('HOLD')},{CL(pc(y))}{op_row}-{R('LN_I')}*{R('MCONST')},0)"
              f"+IF({y}={R('HOLD')},{CL(pc(y))}{rev_row}-({R('LN_I')}*(1+{R('MRATE')})^(12*{R('HOLD')})-({R('LN_I')}*{R('MCONST')}/12)*((1+{R('MRATE')})^(12*{R('HOLD')})-1)/{R('MRATE')}),0)",
              style="calc", fmt=F_ACCT, align="right")
    r += 1
    lh_row = r
    s.put(r, L, "  Levered CF ② (equity)", style="note", align="left")
    s.put(r, ACQ, f"=-{R('EQ_H')}", style="calc", fmt=F_ACCT, align="right")
    for y in range(1, 11):
        s.put(r, pc(y), f"=IF({y}<={R('HOLD')},{CL(pc(y))}{op_row}-{R('LN_H')}*{R('MCONST')},0)"
              f"+IF({y}={R('HOLD')},{CL(pc(y))}{rev_row}-({R('LN_H')}*(1+{R('MRATE')})^(12*{R('HOLD')})-({R('LN_H')}*{R('MCONST')}/12)*((1+{R('MRATE')})^(12*{R('HOLD')})-1)/{R('MRATE')}),0)",
              style="calc", fmt=F_ACCT, align="right")
    r += 1
    two("IRRU_I", "IRRU_H", "Unlevered IRR",
        f"=IRR({CL(ACQ)}{ui_row}:{CL(pc(10))}{ui_row})", f"=IRR({CL(ACQ)}{uh_row}:{CL(pc(10))}{uh_row})", F_PCT1)
    two("EMU_I", "EMU_H", "Unlevered equity multiple",
        f"=SUM({CL(pc(1))}{ui_row}:{CL(pc(10))}{ui_row})/-{CL(ACQ)}{ui_row}",
        f"=SUM({CL(pc(1))}{uh_row}:{CL(pc(10))}{uh_row})/-{CL(ACQ)}{uh_row}", F_MULT)
    two("IRRL_I", "IRRL_H", "Levered IRR",
        f"=IRR({CL(ACQ)}{li_row}:{CL(pc(10))}{li_row})", f"=IRR({CL(ACQ)}{lh_row}:{CL(pc(10))}{lh_row})", F_PCT1)
    two("EML_I", "EML_H", "Levered equity multiple",
        f"=SUM({CL(pc(1))}{li_row}:{CL(pc(10))}{li_row})/{R('EQ_I')}",
        f"=SUM({CL(pc(1))}{lh_row}:{CL(pc(10))}{lh_row})/{R('EQ_H')}", F_MULT)
    two("DSCR_I", "DSCR_H", "Year-1 DSCR",
        f"={R('ASIS_NOI')}/({R('LN_I')}*{R('MCONST')})", f"={R('ASIS_NOI')}/({R('LN_H')}*{R('MCONST')})", F_MULT)
    # cash-on-cash (income scenario) — levered operating CF (pre-reversion) ÷ equity
    coc_row = r
    s.put(r, L, "  Cash-on-cash ① by year (levered)", style="note", align="left")
    for y in range(1, 11):
        s.put(r, pc(y), f"=IF({y}<={R('HOLD')},({CL(pc(y))}{op_row}-{R('LN_I')}*{R('MCONST')})/{R('EQ_I')},0)",
              style="calc", fmt=F_PCT1, align="right")
    r += 1
    two("COC_AVG_I", "COC_AVG_H", "Avg cash-on-cash (Yr 1–hold)",
        f"=SUM({CL(pc(1))}{coc_row}:{CL(pc(10))}{coc_row})/{R('HOLD')}",
        f"=(( {R('SUMOP5')}-{R('HOLD')}*{R('LN_H')}*{R('MCONST')})/{R('HOLD')})/{R('EQ_H')}", F_PCT1)
    r += 1

    # ================= land appreciation break-even =================
    s.section(r, L, 13, "COVERED-LAND JUSTIFICATION  —  land appreciation required to make the HBU price work"); r += 1
    # required exit land value so that levered equity earns target on the HBU price
    # simple frame: additional value needed = HBU price − income value, capitalized as land upside
    der("HBU_GAP", "Premium paid over income value", f"={px_hbu}-{R('PRICE')}", F_ACCT_TOP, "the dirt / optionality cost")
    der("BE_APPREC", "Land CAGR over hold to recover the premium",
        f"=(({R('PRICE')}+{R('HBU_GAP')})/{R('PRICE')})^(1/{R('HOLD')})-1", F_PCT1,
        "annual land-value growth needed just to break even on the premium at exit")
    s.put(r, L, "Read", style="warn", align="left")
    s.put(r, 4, "At the covered-land price the income IRR is thin — the return thesis is land appreciation / redevelopment, "
                "not current yield. The premium only pencils if land compounds at ≥ the rate above over the hold.",
          style="warn", align="left", merge=(r, 13)); r += 2

    # ================= return attribution (value-creation bridge) =================
    s.section(r, L, 13, "RETURN ATTRIBUTION  —  where the unlevered profit comes from  (at income price)"); r += 1
    der("EXIT_FWD", "Exit-year forward NOI (grown into sale)", f"={R('EXITNOI')}*(1+{R('GROW')})", F_ACCT_TOP)
    der("GROSS_EXIT", "Gross exit value (from consolidated reversion)", f"={R('REV_BASE')}/(1-{R('COS')})", F_ACCT_TOP)
    der("V_OPCF", "＋ Operating cash flow (Yr 1–hold)", f"={R('SUMOP5')}", F_ACCT)
    der("V_NOIG", "＋ Value from NOI growth / lease-up", f"=({R('EXIT_FWD')}-{R('ASIS_NOI')})/{R('CAP')}", F_ACCT)
    der("V_CAP", "＋/− Value from cap-rate movement", f"={R('GROSS_EXIT')}-{R('EXIT_FWD')}/{R('CAP')}", F_ACCT)
    der("V_COS", "− Cost of sale at exit", f"=-{R('GROSS_EXIT')}*{R('COS')}", F_ACCT)
    der("V_CLOSE", "− Acquisition transaction costs", f"=-{R('PRICE')}*{R('CLOSE')}", F_ACCT)
    s.put(r, L, "UNLEVERED PROFIT (sum of the above)", style="subtotal", align="left")
    s.put(r, 3, f"={R('V_OPCF')}+{R('V_NOIG')}+{R('V_CAP')}+{R('V_COS')}+{R('V_CLOSE')}",
          style="calc", fmt=F_ACCT_TOP, align="right", bold=True, name="UNLEV_PROFIT")
    r += 1
    s.put(r, L, "Check: Σ unlevered cash flow (income scenario)", style="note", align="left")
    s.put(r, 3, f"=SUM({CL(ACQ)}{ui_row}:{CL(pc(10))}{ui_row})", style="calc", fmt=F_ACCT, align="right")
    s.put(r, 4, "ties to unlevered profit above", style="note", align="left", merge=(r, 13)); r += 2

    # ================= break-even =================
    s.section(r, L, 13, "BREAK-EVEN  —  downside guardrails (at income price)"); r += 1
    der("SUMOP5_LEV", "Σ levered operating CF (Yr 1–hold)", f"={R('SUMOP5')}-{R('HOLD')}*{R('DS')}", F_ACCT_TOP)
    der("BE_EXITCAP", "Break-even exit cap (return of capital, 1.00x)",
        f"={R('EXIT_FWD')}*(1-{R('COS')})/({R('EQ_I')}-{R('SUMOP5_LEV')}+{R('LBAL5')})", F_PCT2,
        "max exit cap before levered equity < 1.0x")
    der("BE_CUSHION", "Cushion vs. underwritten exit cap", f"={R('BE_EXITCAP')}-{R('SCAP')}", F_PCT2, "bps of exit-cap softening tolerable")
    r += 1

    # ================= valuation bridge =================
    s.section(r, L, 13, "VALUATION BRIDGE  —  income basis vs. highest-and-best-use (covered land)"); r += 1
    s.put(r, L, "Income-based purchase price (this tab)", style="label_b", align="left")
    s.put(r, 3, f"={R('PRICE')}", style="calc", fmt=F_ACCT_TOP, align="right", name="PX_INCOME")
    s.put(r, 4, "what the combined cash flows support", style="note", align="left", merge=(r, 13)); r += 1
    s.put(r, L, "HBU covered-land acquisition (Assemblage tab)", style="label_b", align="left")
    s.put(r, 3, f"={px_hbu}", style="calc", color="008000", fmt=F_ACCT_TOP, align="right", name="PX_HBU")
    s.put(r, 4, "land value + 20% assemblage premium to control", style="note", align="left", merge=(r, 13)); r += 1
    s.put(r, L, "PREMIUM OVER INCOME VALUE (dirt + optionality)", style="grand", align="left")
    s.put(r, 3, f"={R('PX_HBU')}-{R('PX_INCOME')}", style="grand", fmt=F_ACCT_TOP, align="right", name="PREMIUM"); r += 1
    s.put(r, L, "Premium as % of income value", style="label", align="left")
    s.put(r, 3, f"={R('PREMIUM')}/{R('PX_INCOME')}", style="calc", fmt=F_PCT1, align="right"); r += 2

    # expose portfolio returns for other tabs (correct, consolidated)
    s.reg["IRR_U"] = s.reg["IRRU_I"]; s.reg["IRR_L"] = s.reg["IRRL_I"]
    s.reg["IRR_U_HBU"] = s.reg["IRRU_H"]; s.reg["IRR_L_HBU"] = s.reg["IRRL_H"]

    # ================= sensitivity 1: value vs cap x NOI =================
    s.section(r, L, 13, "SENSITIVITY ①  —  income value  ($M)  ·  going-in cap (rows) × as-is NOI (cols)"); r += 1
    caps = [0.0600, 0.0625, 0.0650, 0.0675, 0.0700]
    noi_facts = [0.90, 0.95, 1.00, 1.05, 1.10]
    s.put(r, L, "Cap  ╲  NOI", style="subhead", align="center")
    for j, f in enumerate(noi_facts):
        s.put(r, 4 + j, f, style="input", fmt="0%", align="center")
    r += 1
    s1_top = r
    for cap in caps:
        s.put(r, L, cap, style="input", fmt=F_PCT2, align="center")
        for j in range(len(noi_facts)):
            s.put(r, 4 + j, f"={R('ASIS_NOI')}*{CL(4+j)}{s1_top-1}/{CL(L)}{r}/1000000",
                  style="calc", fmt="#,##0.0", align="center")
        r += 1
    s.ws.conditional_formatting.add(f"D{s1_top}:H{r-1}", HEAT)
    r += 1

    # ================= sensitivity 2: levered EM vs price x exit cap =================
    s.section(r, L, 13, "SENSITIVITY ②  —  levered equity multiple  ·  purchase price (rows) × exit cap (cols)"); r += 1
    prices = [65000000, 75000000, 85000000, 95000000, 102000000]
    xcaps = [0.0600, 0.0650, 0.0700, 0.0750, 0.0800]
    s.put(r, L, "Price ╲ exit cap", style="subhead", align="center")
    for j, xc in enumerate(xcaps):
        s.put(r, 4 + j, xc, style="input", fmt=F_PCT2, align="center")
    r += 1
    s2_top = r
    for p in prices:
        s.put(r, L, p, style="input", fmt=F_ACCT, align="center")
        for j in range(len(xcaps)):
            xc = f"{CL(4+j)}{s2_top-1}"
            P = f"{CL(L)}{r}"
            # closed-form levered EM: (Σop5 − hold×DS(P) + reversion(xc) − loanbal5(P)) / equity(P)
            loanP = f"MIN({P}*{R('LTV')},{R('L_DSCR')},{R('L_DY')})"
            dsP = f"({loanP}*{R('MCONST')})"
            bal = f"({loanP}*(1+{R('MRATE')})^(12*{R('HOLD')})-({dsP}/12)*((1+{R('MRATE')})^(12*{R('HOLD')})-1)/{R('MRATE')})"
            eqP = f"({P}*(1+{R('CLOSE')})-{loanP})"
            rev = f"({R('EXITNOI')}*(1+{R('GROW')})/{xc}*(1-{R('COS')}))"
            s.put(r, 4 + j,
                  f"=({R('SUMOP5')}-{R('HOLD')}*{dsP}+{rev}-{bal})/{eqP}",
                  style="calc", fmt=F_MULT, align="center")
        r += 1
    s.ws.conditional_formatting.add(f"D{s2_top}:H{r-1}", HEAT)
    r += 1
    s.put(r, L, "Note", style="note", align="left")
    s.put(r, 4, "Sensitivity ② uses a blended exit cap on consolidated exit-year NOI and lesser-of debt sizing at each price; "
                "operating cash flows are held at the underwritten level. Green = higher multiple.",
          style="note", align="left", merge=(r, 13)); r += 1

    # row anchor for the Scenarios tab
    s.reg["ROW_OPCF"] = f"D{op_row}"

    s.freeze("C6")
    return s
