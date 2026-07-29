"""
tab_asset.py — generic income-producing asset tab (stabilized / NNN).
Mirrors the validated Shahidi cash-flow engine but parameterized by a config.
Produces: facts, lease/rent detail, blue assumptions, 10-yr DCF, reversion,
debt, unlevered + levered returns (IRR/EM/NPV), and a P&L. Value is driven off
the rent the asset produces (income approach) with a land-basis cross-check.
"""
from openpyxl.utils import get_column_letter
from mblib import (F_ACCT, F_ACCT_TOP, F_PCT1, F_PCT2, F_MULT, F_PSF, F_NUM, F_NUM2, F_YR)

L = 2
ACQ = 3
def pc(y): return 3 + y
def CL(c): return get_column_letter(c)


def build(s, cfg, amap=None):
    amap = amap or {}
    a = cfg["inp"]
    s.colw({"A": 2.5, "B": 34, "C": 13, "D": 12, "E": 12, "F": 12, "G": 12,
            "H": 12, "I": 12, "J": 12, "K": 12, "L": 12, "M": 12})
    r = 1
    # banner
    s.put(r, L, cfg["title"], style="banner", align="left", merge=(r, 13)); s.rowh(r, 26); r += 1
    s.put(r, L, cfg["subtitle"], style="banner_sub", align="left", merge=(r, 13)); s.rowh(r, 18); r += 2

    # facts
    s.section(r, L, 13, cfg["facts_header"]); r += 1
    for lab, val, verified in cfg["facts"]:
        st = "verified" if verified else "calc"
        s.put(r, L, lab, style=st, align="left")
        s.put(r, 4, val, style=st, align="left", merge=(r, 13)); r += 1
    r += 1

    # lease / rent detail
    if cfg.get("leases"):
        s.section(r, L, 13, cfg["lease_header"]); r += 1
        hdr = ["Tenant / component", "Category", "SF", "$/SF NNN", "Annual rent", "Note"]
        cols = [L, 5, 7, 8, 9, 12]
        spans = {2: 4, 5: 6, 9: 11, 12: 13}
        for h, c in zip(hdr, cols):
            endc = spans.get(c, c)
            s.put(r, c, h, style="subhead", align="left" if c in (L, 5, 12) else "center",
                  merge=(r, endc) if endc != c else None)
        r += 1
        lstart = r
        for ten, cat, sf, rent, note in cfg["leases"]:
            s.put(r, L, ten, style="calc", align="left", merge=(r, 4))
            s.put(r, 5, cat, style="calc", align="left", merge=(r, 6))
            s.put(r, 7, sf, style="calc", fmt=F_NUM, align="right")
            s.put(r, 8, rent, style="input", fmt=F_PSF, align="right")
            s.put(r, 9, f"={CL(7)}{r}*{CL(8)}{r}", style="calc", fmt=F_ACCT, align="right", merge=(r, 11))
            s.put(r, 12, note, style="calc", align="left", merge=(r, 13))
            r += 1
        if cfg.get("vacant_sf"):
            s.put(r, L, "VACANT / available", style="warn", align="left", merge=(r, 4))
            s.put(r, 5, "Lease-up upside", style="warn", align="left", merge=(r, 6))
            s.put(r, 7, cfg["vacant_sf"], style="warn", fmt=F_NUM, align="right")
            s.put(r, 8, "—", style="warn", align="right")
            s.put(r, 9, "—", style="warn", align="right", merge=(r, 11))
            s.put(r, 12, "value-add", style="warn", align="left", merge=(r, 13))
            r += 1
        lend = r - 1
        s.put(r, L, "TOTAL / WTD AVG", style="total", align="left", merge=(r, 4))
        s.put(r, 5, "", style="total", merge=(r, 6))
        s.put(r, 7, f"=SUM({CL(7)}{lstart}:{CL(7)}{lend})", style="total", fmt=F_NUM, align="right")
        occ_end = lend - (1 if cfg.get("vacant_sf") else 0)
        s.put(r, 8, f"=IFERROR(SUM({CL(9)}{lstart}:{CL(9)}{occ_end})/SUM({CL(7)}{lstart}:{CL(7)}{occ_end}),0)",
              style="total", fmt=F_PSF, align="right")
        s.put(r, 9, f"=SUM({CL(9)}{lstart}:{CL(9)}{occ_end})", style="total", fmt=F_ACCT_TOP, align="right", merge=(r, 11))
        s.put(r, 12, cfg.get("lease_note", ""), style="total", align="left", merge=(r, 13))
        r += 2

    # assumptions
    s.section(r, L, 13, "ASSUMPTIONS  —  🔵 blue = hardcoded inputs"); r += 1
    def grp(t):
        nonlocal r
        s.put(r, L, t, style="subhead", align="left", merge=(r, 13)); r += 1
    def inp(name, label, key, fmt, note):
        nonlocal r
        s.put(r, L, label, style="label", align="left")
        if name in amap:
            s.put(r, 3, f"='Assumptions'!{amap[name]}", style="calc", color="008000", fmt=fmt, align="right", name=name)
            s.put(r, 4, "🟢 Assumptions", style="note", align="left", merge=(r, 13))
        else:
            s.put(r, 3, a[key], style="input", fmt=fmt, align="right", name=name)
            s.put(r, 4, note, style="note", align="left", merge=(r, 13))
        r += 1

    grp("Property & revenue")
    inp("GLA", "Rentable SF", "gla", F_NUM, cfg["src"]["gla"])
    inp("GLANDSF", "Land (SF)", "land_sf", F_NUM, cfg["src"].get("land", "🔶 parcel land"))
    inp("OCC0", "In-place occupancy", "occ0", F_PCT1, cfg["src"]["occ0"])
    inp("STABOCC", "Stabilized occupancy", "stab_occ", F_PCT1, cfg["src"]["stab_occ"])
    if cfg.get("mrent_from_slb"):
        AL = lambda n: f"'Assumptions'!{amap[n]}"
        s.put(r, L, "Blended NNN rent ($/SF) — leaseback + pad", style="label", align="left")
        s.put(r, 3, f"=({AL('SLB_SF')}*{AL('SLB_RENT')}+{AL('SBUX_SF')}*{AL('SBUX_RENT')})/{s.reg['GLA']}",
              style="calc", color="008000", fmt=F_PSF, align="right", name="MRENT")
        s.put(r, 4, "🟢 computed from Publix leaseback + Starbucks pad (Assumptions)", style="note", align="left", merge=(r, 13)); r += 1
    else:
        inp("MRENT", "Market rent ($/SF NNN)", "market_rent", F_PSF, cfg["src"]["rent"])
    inp("RGROW", "Rent growth", "rent_growth", F_PCT1, "🔶")
    inp("CLOSS", "Credit & collection loss", "credit_loss", F_PCT1, "🔶")
    inp("MILL", "Effective millage", "millage", F_PCT2, cfg["src"]["mill"])
    grp("Operating expenses")
    inp("INS", "Insurance ($/yr)", "insurance", F_ACCT_TOP, "🔶 AE flood")
    inp("CAM", "CAM ($/yr, recoverable)", "cam", F_ACCT_TOP, "🔶")
    inp("RM", "Repairs & maintenance ($/yr)", "rm", F_ACCT_TOP, "🔶 non-recoverable")
    inp("MGMT", "Management fee (% EGR)", "mgmt_pct", F_PCT1, "🔶")
    inp("EGROW", "Expense growth", "exp_growth", F_PCT1, "🔶")
    grp("Capital")
    inp("TI", "Tenant improvements ($/SF)", "ti_psf", F_PSF, "🔶 on leased-up SF")
    inp("LC", "Leasing commissions ($/SF)", "lc_psf", F_PSF, "🔶 on leased-up SF")
    inp("RES", "Replacement reserves ($/SF/yr)", "reserve_psf", F_PSF, "🔶")
    inp("ROLL", "Rollover leasing reserve ($/SF/yr)", "rollover_psf", F_PSF, "🔶")
    grp("Acquisition & financing")
    inp("PRICE", "Purchase price / basis", "price", F_ACCT_TOP, cfg["src"]["price"])
    inp("CLOSE", "Closing & acq costs (% price)", "closing_pct", F_PCT1, "🔶")
    inp("LTV", "Senior loan LTV", "ltv", F_PCT1, "🔶")
    inp("RATE", "Senior loan rate", "loan_rate", F_PCT2, "🔶 mid-2026")
    inp("AMORT", "Senior loan amortization (yrs)", "amort", F_YR, "🔶")
    grp("Valuation & hold")
    inp("GICAP", "Going-in cap", "goingin_cap", F_PCT2, cfg["src"]["cap"])
    inp("EXITCAP", "Exit / reversion cap", "exit_cap", F_PCT2, cfg["src"]["exit"])
    inp("COS", "Cost of sale at exit", "cost_sale", F_PCT1, "🔶")
    inp("HOLD", "Hold period (yrs)", "hold", F_YR, "🔶")
    inp("DISC", "Discount rate (NPV)", "disc", F_PCT1, "🔶")
    s.put(r, L, "Land value ($/SF)", style="label", align="left")
    if "GLAND_PSF" in amap:
        s.put(r, 3, f"='Assumptions'!{amap['GLAND_PSF']}", style="calc", color="008000", fmt=F_PSF, align="right", name="GLAND_PSF")
        s.put(r, 4, "🟢 Assumptions", style="note", align="left", merge=(r, 13))
    else:
        s.put(r, 3, cfg.get("land_psf", 248), style="input", fmt=F_PSF, align="right", name="GLAND_PSF")
        s.put(r, 4, "🔶 corridor land comp", style="note", align="left", merge=(r, 13))
    r += 1
    r += 1

    g = s.reg
    def R(n): return g[n]

    # derived / debt
    s.section(r, L, 13, "DERIVED VALUES  &  DEBT SIZING"); r += 1
    def der(name, label, formula, fmt, note=""):
        nonlocal r
        s.put(r, L, label, style="label", align="left")
        s.put(r, 3, formula, style="calc", fmt=fmt, align="right", name=name)
        if note: s.put(r, 4, note, style="note", align="left", merge=(r, 13))
        r += 1
    der("LOAN", "Senior loan amount", f"={R('PRICE')}*{R('LTV')}", F_ACCT_TOP, "price × LTV")
    der("MRATE", "Monthly rate", f"={R('RATE')}/12", F_PCT2)
    der("DS", "Annual debt service",
        f"={R('LOAN')}*({R('MRATE')}/(1-(1+{R('MRATE')})^(-{R('AMORT')}*12)))*12", F_ACCT_TOP)
    der("EQ_ACQ", "Equity at acquisition", f"={R('PRICE')}*(1+{R('CLOSE')})-{R('LOAN')}", F_ACCT_TOP)
    der("ABSORB", "Absorption SF (lease-up)", f"={R('GLA')}*({R('STABOCC')}-{R('OCC0')})", F_NUM)
    der("INIT_LEASE", "Initial leasing capital", f"={R('ABSORB')}*({R('TI')}+{R('LC')})", F_ACCT_TOP)
    der("TCB", "Total cost basis", f"={R('PRICE')}*(1+{R('CLOSE')})+{R('INIT_LEASE')}", F_ACCT_TOP)
    der("EQ_REQ", "Total equity required", f"={R('TCB')}-{R('LOAN')}", F_ACCT_TOP)
    der("INCVAL", "Income value (Yr-1 NOI ÷ going-in cap)", "=0", F_ACCT_TOP, "computed below")
    incval_cell = R("INCVAL")
    der("LANDVAL", "Land value (land SF × $/SF)",
        f"={R('GLANDSF')}*{R('GLAND_PSF')}", F_ACCT_TOP,
        "land SF × land $/SF (both blue inputs above)")
    r += 1

    # ---- cash flow ----
    s.section(r, L, 13, "OPERATING CASH FLOW  —  10-year annual"); r += 1
    s.put(r, L, "Year", style="subhead", align="left")
    s.put(r, ACQ, "0 / Acq", style="subhead", align="center")
    for y in range(1, 11):
        s.put(r, pc(y), y, style="subhead", align="center", fmt=F_YR)
    r += 1
    s.put(r, L, "Year ending", style="note", align="left")
    s.put(r, ACQ, "At close", style="note", align="center")
    for y in range(1, 11):
        s.put(r, pc(y), f"Dec-{2026+y}", style="note", align="center")
    r += 1

    rows = {}
    def cfrow(name, label, bold=False, top=False):
        nonlocal r
        rows[name] = r
        st = "subtotal" if top else ("label_b" if bold else "label")
        s.put(r, L, label, style=st, align="left")
        return r

    ro = cfrow("OCC", "Occupancy")
    for y in range(1, 11):
        f = (f"=({R('OCC0')}+{R('STABOCC')})/2" if y == 1 else f"={R('STABOCC')}")
        s.put(ro, pc(y), f, fmt=F_PCT1, align="right")
    r += 1
    rl = cfrow("LEASED", "Leased SF")
    for y in range(1, 11):
        s.put(rl, pc(y), f"={R('GLA')}*{CL(pc(y))}{rows['OCC']}", fmt=F_NUM, align="right")
    r += 1
    s.put(r, L, "REVENUE", style="subhead", align="left", merge=(r, 13)); r += 1
    slb_term = cfg.get("slb_term") and "TERM" in amap
    term_ref = f"'Assumptions'!{amap['TERM']}" if slb_term else None
    rb = cfrow("BASE", "Base rental income" + (" (leaseback cliffs at term)" if slb_term else ""))
    for y in range(1, 11):
        core = f"{CL(pc(y))}{rows['LEASED']}*{R('MRENT')}*(1+{R('RGROW')})^({y}-1)"
        f = f"=IF({y}<={term_ref},{core},0)" if slb_term else f"={core}"
        s.put(rb, pc(y), f, fmt=(F_ACCT_TOP if y == 1 else F_ACCT), align="right")
    r += 1
    rreco = cfrow("RECOV", "Recoverable opex (gross)")
    for y in range(1, 11):
        s.put(rreco, pc(y), f"=({R('PRICE')}*{R('MILL')}+{R('INS')}+{R('CAM')})*(1+{R('EGROW')})^({y}-1)",
              fmt=F_ACCT, align="right")
    r += 1
    rreim = cfrow("REIMB", "Expense reimbursements (NNN)")
    for y in range(1, 11):
        s.put(rreim, pc(y), f"={CL(pc(y))}{rows['RECOV']}*{CL(pc(y))}{rows['OCC']}", fmt=F_ACCT, align="right")
    r += 1
    rcl = cfrow("CLOSSX", "Less: credit & collection loss")
    for y in range(1, 11):
        s.put(rcl, pc(y), f"=-{CL(pc(y))}{rows['BASE']}*{R('CLOSS')}", fmt=F_ACCT, align="right")
    r += 1
    regr = cfrow("EGR", "Effective Gross Revenue", top=True)
    for y in range(1, 11):
        s.put(regr, pc(y), f"={CL(pc(y))}{rows['BASE']}+{CL(pc(y))}{rows['REIMB']}+{CL(pc(y))}{rows['CLOSSX']}",
              fmt=(F_ACCT_TOP if y == 1 else F_ACCT), align="right")
    r += 1
    s.put(r, L, "OPERATING EXPENSES", style="subhead", align="left", merge=(r, 13)); r += 1
    rtax = cfrow("TAX", "Real-estate taxes (reassessed)")
    for y in range(1, 11):
        s.put(rtax, pc(y), f"={R('PRICE')}*{R('MILL')}*(1+{R('EGROW')})^({y}-1)", fmt=F_ACCT, align="right")
    r += 1
    rins = cfrow("INSX", "Insurance")
    for y in range(1, 11):
        s.put(rins, pc(y), f"={R('INS')}*(1+{R('EGROW')})^({y}-1)", fmt=F_ACCT, align="right")
    r += 1
    rcam = cfrow("CAMX", "CAM")
    for y in range(1, 11):
        s.put(rcam, pc(y), f"={R('CAM')}*(1+{R('EGROW')})^({y}-1)", fmt=F_ACCT, align="right")
    r += 1
    rrm = cfrow("RMX", "Repairs & maintenance")
    for y in range(1, 11):
        s.put(rrm, pc(y), f"={R('RM')}*(1+{R('EGROW')})^({y}-1)", fmt=F_ACCT, align="right")
    r += 1
    rmg = cfrow("MGX", "Management fee")
    for y in range(1, 11):
        s.put(rmg, pc(y), f"={CL(pc(y))}{rows['EGR']}*{R('MGMT')}", fmt=F_ACCT, align="right")
    r += 1
    rtox = cfrow("TOPEX", "Total operating expenses", top=True)
    for y in range(1, 11):
        s.put(rtox, pc(y), f"=SUM({CL(pc(y))}{rows['TAX']}:{CL(pc(y))}{rows['MGX']})",
              fmt=(F_ACCT_TOP if y == 1 else F_ACCT), align="right")
    r += 1
    rnoi = cfrow("NOI", "NET OPERATING INCOME", bold=True, top=True)
    for y in range(1, 11):
        s.put(rnoi, pc(y), f"={CL(pc(y))}{rows['EGR']}-{CL(pc(y))}{rows['TOPEX']}",
              fmt=(F_ACCT_TOP if y == 1 else F_ACCT), align="right", bold=True)
    r += 1
    s.put(r, L, "CAPITAL", style="subhead", align="left", merge=(r, 13)); r += 1
    rres = cfrow("RESV", "Replacement reserves")
    for y in range(1, 11):
        s.put(rres, pc(y), f"={R('GLA')}*{R('RES')}*(1+{R('EGROW')})^({y}-1)", fmt=F_ACCT, align="right")
    r += 1
    rlease = cfrow("LEASECAP", "Leasing costs")
    for y in range(1, 11):
        if y == 1:
            f = f"={R('ABSORB')}*({R('TI')}+{R('LC')})+{R('GLA')}*{R('ROLL')}"
        else:
            f = f"={R('GLA')}*{R('ROLL')}*(1+{R('EGROW')})^({y}-1)"
        s.put(rlease, pc(y), f, fmt=F_ACCT, align="right")
    r += 1
    rtcap = cfrow("TCAP", "Total capital", top=True)
    for y in range(1, 11):
        s.put(rtcap, pc(y), f"={CL(pc(y))}{rows['RESV']}+{CL(pc(y))}{rows['LEASECAP']}",
              fmt=(F_ACCT_TOP if y == 1 else F_ACCT), align="right")
    r += 1
    runcf = cfrow("UNCF", "Unlevered cash flow (before reversion)", bold=True, top=True)
    s.put(runcf, ACQ, f"=-{R('PRICE')}*(1+{R('CLOSE')})", fmt=F_ACCT_TOP, align="right", bold=True)
    for y in range(1, 11):
        s.put(runcf, pc(y), f"={CL(pc(y))}{rows['NOI']}-{CL(pc(y))}{rows['TCAP']}",
              fmt=(F_ACCT_TOP if y == 1 else F_ACCT), align="right", bold=True)
    r += 2

    # reversion & levered
    s.section(r, L, 13, "REVERSION, DEBT SERVICE  &  LEVERED CASH FLOW"); r += 1
    exit_land = cfg.get("exit_land")
    rns = cfrow("NETSALE", "Net sale — land value (covered-land exit)" if exit_land else "Net sale proceeds (exit year)")
    for y in range(1, 11):
        if exit_land:
            f = f"=IF({y}={R('HOLD')},{R('LANDVAL')}*(1-{R('COS')}),0)"
        else:
            f = f"=IF({y}={R('HOLD')},({CL(pc(y))}{rows['NOI']}*(1+{R('RGROW')})/{R('EXITCAP')})*(1-{R('COS')}),0)"
        s.put(rns, pc(y), f, fmt=F_ACCT, align="right")
    r += 1
    rupcf = cfrow("PROJCF", "UNLEVERED PROJECT CASH FLOW", bold=True, top=True)
    s.put(rupcf, ACQ, f"={CL(ACQ)}{rows['UNCF']}", fmt=F_ACCT_TOP, align="right", bold=True)
    for y in range(1, 11):
        s.put(rupcf, pc(y),
              f"=IF({y}<={R('HOLD')},{CL(pc(y))}{rows['UNCF']},0)+{CL(pc(y))}{rows['NETSALE']}",
              fmt=(F_ACCT_TOP if y == 1 else F_ACCT), align="right", bold=True)
    r += 1
    rds = cfrow("DSV", "Debt service")
    for y in range(1, 11):
        s.put(rds, pc(y), f"=IF({y}<={R('HOLD')},-{R('DS')},0)", fmt=F_ACCT, align="right")
    r += 1
    rlb = cfrow("LBAL", "Loan balance (year-end)")
    for y in range(1, 11):
        s.put(rlb, pc(y),
              f"={R('LOAN')}*(1+{R('MRATE')})^(12*{y})-({R('DS')}/12)*((1+{R('MRATE')})^(12*{y})-1)/{R('MRATE')}",
              fmt=F_ACCT, align="right")
    r += 1
    rdscr = cfrow("DSCR", "DSCR (NOI ÷ DS)")
    for y in range(1, 11):
        s.put(rdscr, pc(y), f"={CL(pc(y))}{rows['NOI']}/{R('DS')}", fmt=F_MULT, align="right")
    r += 1
    rdy = cfrow("DYLD", "Debt yield (NOI ÷ loan)")
    for y in range(1, 11):
        s.put(rdy, pc(y), f"={CL(pc(y))}{rows['NOI']}/{R('LOAN')}", fmt=F_PCT1, align="right")
    r += 1
    rlev = cfrow("LEVCF", "LEVERED CASH FLOW (equity)", bold=True, top=True)
    s.put(rlev, ACQ, f"=-{R('EQ_ACQ')}", fmt=F_ACCT_TOP, align="right", bold=True)
    for y in range(1, 11):
        s.put(rlev, pc(y),
              f"=IF({y}<={R('HOLD')},{CL(pc(y))}{rows['UNCF']}-{R('DS')},0)"
              f"+IF({y}={R('HOLD')},{CL(pc(y))}{rows['NETSALE']}-{CL(pc(y))}{rows['LBAL']},0)",
              fmt=(F_ACCT_TOP if y == 1 else F_ACCT), align="right", bold=True)
    r += 2

    # patch income value now that NOI row exists
    s.ws[incval_cell] = f"={CL(pc(1))}{rows['NOI']}/{R('GICAP')}"

    # returns
    s.section(r, L, 13, "RETURNS  —  unlevered & levered"); r += 1
    proj = f"{CL(ACQ)}{rows['PROJCF']}:{CL(pc(10))}{rows['PROJCF']}"
    lev = f"{CL(ACQ)}{rows['LEVCF']}:{CL(pc(10))}{rows['LEVCF']}"
    proj1 = f"{CL(pc(1))}{rows['PROJCF']}:{CL(pc(10))}{rows['PROJCF']}"
    lev1 = f"{CL(pc(1))}{rows['LEVCF']}:{CL(pc(10))}{rows['LEVCF']}"
    def ret(name, label, formula, fmt):
        nonlocal r
        s.put(r, L, label, style="label", align="left")
        s.put(r, 3, formula, style="calc", fmt=fmt, align="right", name=name)
        r += 1
    ret("IRR_U", "Unlevered IRR", f'=IFERROR(IRR({proj}),"n/m")', F_PCT1)
    ret("EM_U", "Unlevered equity multiple", f"=SUM({proj1})/-{CL(ACQ)}{rows['PROJCF']}", F_MULT)
    ret("NPV_U", "Unlevered NPV @ discount rate", f"={CL(ACQ)}{rows['PROJCF']}+NPV({R('DISC')},{proj1})", F_ACCT_TOP)
    ret("IRR_L", "Levered IRR", f'=IFERROR(IRR({lev}),"n/m")', F_PCT1)
    ret("EM_L", "Levered equity multiple", f"=SUM({lev1})/{R('EQ_ACQ')}", F_MULT)
    ret("NOI1", "Year-1 NOI (pro-forma)", f"={CL(pc(1))}{rows['NOI']}", F_ACCT_TOP)
    ret("AS_IS_NOI", "As-is in-place NOI (at occ0, no lease-up)",
        f"=({R('GLA')}*{R('OCC0')}*{R('MRENT')})"
        f"+({R('PRICE')}*{R('MILL')}+{R('INS')}+{R('CAM')})*{R('OCC0')}"
        f"-({R('GLA')}*{R('OCC0')}*{R('MRENT')})*{R('CLOSS')}"
        f"-({R('PRICE')}*{R('MILL')}+{R('INS')}+{R('CAM')}+{R('RM')})"
        f"-(({R('GLA')}*{R('OCC0')}*{R('MRENT')})+({R('PRICE')}*{R('MILL')}+{R('INS')}+{R('CAM')})*{R('OCC0')}"
        f"-({R('GLA')}*{R('OCC0')}*{R('MRENT')})*{R('CLOSS')})*{R('MGMT')}", F_ACCT_TOP)
    ret("GOINGIN", "Going-in cap (as-is NOI ÷ price)", f"={R('AS_IS_NOI')}/{R('PRICE')}", F_PCT2)
    ret("STABNOI", "Stabilized NOI (Yr 2)", f"={CL(pc(2))}{rows['NOI']}", F_ACCT_TOP)
    ret("STABVAL", "Stabilized value (÷ exit cap)", f"={R('STABNOI')}/{R('EXITCAP')}", F_ACCT_TOP)
    ret("YOC", "Stabilized yield on cost", f"={R('STABNOI')}/{R('TCB')}", F_PCT1)
    ret("DSCR1", "Year-1 DSCR", f"={CL(pc(1))}{rows['NOI']}/{R('DS')}", F_MULT)
    ret("DY1", "Year-1 debt yield", f"={CL(pc(1))}{rows['NOI']}/{R('LOAN')}", F_PCT1)
    r += 1

    # ---- sale-leaseback rent solver ----
    if cfg.get("slb_solver"):
        AL = lambda n: f"'Assumptions'!{amap[n]}"
        s.section(r, L, 13, "SALE-LEASEBACK SOLVER  —  required Publix leaseback rent to hit a target"); r += 1
        s.put(r, L, "Target Year-1 DSCR", style="label", align="left")
        s.put(r, 3, 1.20, style="input", fmt=F_MULT, align="right", name="SLV_TDSCR")
        s.put(r, 4, "🔵 your coverage target", style="note", align="left", merge=(r, 13)); r += 1
        s.put(r, L, "Target going-in cap", style="label", align="left")
        s.put(r, 3, 0.050, style="input", fmt=F_PCT2, align="right", name="SLV_TCAP")
        s.put(r, 4, "🔵 your yield target", style="note", align="left", merge=(r, 13)); r += 1
        s.put(r, L, "K = (1−credit)(1−mgmt)", style="label", align="left")
        s.put(r, 3, f"=(1-{R('CLOSS')})*(1-{R('MGMT')})", style="calc", fmt=F_NUM2, align="right", name="SLV_K"); r += 1
        s.put(r, L, "Recoverable opex (taxes+ins+CAM)", style="label", align="left")
        s.put(r, 3, f"={R('PRICE')}*{R('MILL')}+{R('INS')}+{R('CAM')}", style="calc", fmt=F_ACCT, align="right", name="SLV_RECOV"); r += 1
        base_d = f"(({R('SLV_TDSCR')}*{R('DS')}+{R('SLV_RECOV')}*{R('MGMT')}+{R('RM')})/{R('SLV_K')})"
        base_c = f"(({R('SLV_TCAP')}*{R('PRICE')}+{R('SLV_RECOV')}*{R('MGMT')}+{R('RM')})/{R('SLV_K')})"
        s.put(r, L, "→ Required leaseback rent for target DSCR", style="subtotal", align="left")
        s.put(r, 3, f"=({base_d}-{AL('SBUX_SF')}*{AL('SBUX_RENT')})/{AL('SLB_SF')}", style="calc", fmt=F_PSF, align="right", bold=True, name="SLV_RENT_D")
        s.put(r, 4, "$/SF NNN — set Publix leaseback rent (Assumptions) to this", style="note", align="left", merge=(r, 13)); r += 1
        s.put(r, L, "→ Required leaseback rent for target cap", style="subtotal", align="left")
        s.put(r, 3, f"=({base_c}-{AL('SBUX_SF')}*{AL('SBUX_RENT')})/{AL('SLB_SF')}", style="calc", fmt=F_PSF, align="right", bold=True, name="SLV_RENT_C")
        s.put(r, 4, "$/SF NNN — set Publix leaseback rent (Assumptions) to this", style="note", align="left", merge=(r, 13)); r += 1
        r += 1

    # ---- Publix: leaseback vs dark (what if Publix won't lease back) ----
    if cfg.get("dark_case"):
        AL = lambda n: f"'Assumptions'!{amap[n]}"
        s.section(r, L, 13, "PUBLIX: LEASEBACK vs DARK  —  the deal-critical question: what if Publix won't lease back"); r += 1
        s.put(r, L, "Case", style="subhead", align="left")
        s.put(r, 5, "Leaseback", style="subhead", align="center", merge=(r, 6))
        s.put(r, 7, "Dark (Publix out)", style="subhead", align="center", merge=(r, 8))
        s.put(r, 9, "note", style="subhead", align="left", merge=(r, 13)); r += 1
        # dark case: Publix vacates; only the Starbucks pad remains (separate tenant), NNN on its share
        sbux = f"({AL('SBUX_SF')}*{AL('SBUX_RENT')})"
        recov = f"({R('PRICE')}*{R('MILL')}+{R('INS')}+{R('CAM')})"
        occd = f"({AL('SBUX_SF')}/{R('GLA')})"
        egrd = f"({sbux}+{recov}*{occd}-{sbux}*{R('CLOSS')})"
        opex = f"({R('PRICE')}*{R('MILL')}+{R('INS')}+{R('CAM')}+{R('RM')})"
        noid = f"({egrd}-{opex}-{egrd}*{R('MGMT')})"
        s.put(r, L, "Year-1 NOI", style="label", align="left")
        s.put(r, 5, f"={R('NOI1')}", style="calc", color="008000", fmt=F_ACCT_TOP, align="right", merge=(r, 6))
        s.put(r, 7, f"={noid}", style="calc", fmt=F_ACCT_TOP, align="right", name="NOI_DARK", merge=(r, 8))
        s.put(r, 9, "dark = Starbucks pad only; you carry Publix's taxes & insurance", style="note", align="left", merge=(r, 13)); r += 1
        s.put(r, L, "Annual carry swing (leaseback − dark)", style="label", align="left")
        s.put(r, 5, f"={R('NOI1')}-{R('NOI_DARK')}", style="calc", fmt=F_ACCT, align="right", name="DARK_SWING", merge=(r, 6))
        s.put(r, 7, "per year of vacancy", style="note", align="left", merge=(r, 13)); r += 1
        s.put(r, L, "Entitlement term (yrs)", style="label", align="left")
        s.put(r, 5, f"={AL('TERM')}", style="calc", color="008000", fmt=F_YR, align="center", merge=(r, 6)); r += 1
        s.put(r, L, "→ Extra equity to carry the dark center", style="subtotal", align="left")
        s.put(r, 5, f"={R('DARK_SWING')}*{AL('TERM')}", style="calc", fmt=F_ACCT_TOP, align="right", bold=True, name="DARK_CARRY", merge=(r, 6))
        s.put(r, 7, "cumulative over the entitlement term — funded by equity", style="warn", align="left", merge=(r, 13)); r += 1
        s.put(r, L, "Deal note", style="warn", align="left")
        s.put(r, 4, "The sale-leaseback is DEAL-CRITICAL — Publix is a fee owner, not a seller. Confirm leaseback appetite & rent BEFORE hard money. "
                    "If it goes dark you fund the carry above and lean entirely on land value; the covered-land exit is unchanged.",
              style="warn", align="left", merge=(r, 13)); r += 2

    # P&L
    s.section(r, L, 13, "PROFIT & LOSS STATEMENT  —  Year 1 vs. Stabilized (Yr 2)"); r += 1
    s.put(r, L, "$ / year", style="subhead", align="left")
    s.put(r, 6, "Year 1", style="subhead", align="right", merge=(r, 7))
    s.put(r, 8, "Stabilized", style="subhead", align="right", merge=(r, 9))
    s.put(r, 10, "Per SF (stab.)", style="subhead", align="right", merge=(r, 13)); r += 1
    y1, y2 = pc(1), pc(2)
    def pl(label, rn, sign=1, bold=False, top=False):
        nonlocal r
        st = "subtotal" if top else ("label_b" if bold else "label")
        s.put(r, L, label, style=st, align="left")
        s.put(r, 6, f"={'-' if sign<0 else ''}{CL(y1)}{rows[rn]}", fmt=F_ACCT, align="right", bold=bold, merge=(r, 7))
        s.put(r, 8, f"={'-' if sign<0 else ''}{CL(y2)}{rows[rn]}", fmt=(F_ACCT_TOP if top or bold else F_ACCT), align="right", bold=bold, merge=(r, 9))
        s.put(r, 10, f"={CL(8)}{r}/{R('GLA')}", fmt=F_PSF, align="right", bold=bold, merge=(r, 13))
        r += 1
    pl("Base rental income", "BASE", bold=True)
    pl("Expense reimbursements", "REIMB")
    pl("Less: credit loss", "CLOSSX")
    pl("Effective Gross Income", "EGR", bold=True, top=True)
    pl("Real-estate taxes", "TAX", sign=-1)
    pl("Insurance", "INSX", sign=-1)
    pl("CAM", "CAMX", sign=-1)
    pl("Repairs & maintenance", "RMX", sign=-1)
    pl("Management fee", "MGX", sign=-1)
    pl("Net Operating Income", "NOI", bold=True, top=True)
    pl("Replacement reserves", "RESV", sign=-1)
    pl("Leasing costs", "LEASECAP", sign=-1)
    pl("Cash Flow Before Debt", "UNCF", bold=True, top=True)
    s.put(r, L, "Debt service", style="label", align="left")
    s.put(r, 6, f"=-{R('DS')}", fmt=F_ACCT, align="right", merge=(r, 7))
    s.put(r, 8, f"=-{R('DS')}", fmt=F_ACCT, align="right", merge=(r, 9))
    s.put(r, 10, f"={CL(8)}{r}/{R('GLA')}", fmt=F_PSF, align="right", merge=(r, 13)); r += 1
    cfad = r
    s.put(r, L, "Cash Flow After Debt (levered)", style="subtotal", align="left")
    s.put(r, 6, f"={CL(6)}{cfad-2}-{R('DS')}", fmt=F_ACCT_TOP, align="right", bold=True, merge=(r, 7))
    s.put(r, 8, f"={CL(8)}{cfad-2}-{R('DS')}", fmt=F_ACCT_TOP, align="right", bold=True, merge=(r, 9))
    s.put(r, 10, f"={CL(8)}{r}/{R('GLA')}", fmt=F_PSF, align="right", bold=True, merge=(r, 13)); r += 1

    # row anchors for the consolidated income tab
    s.reg["ROW_NOI"] = f"D{rows['NOI']}"
    s.reg["ROW_CAP"] = f"D{rows['TCAP']}"
    s.reg["ROW_REV"] = f"D{rows['NETSALE']}"

    s.freeze("C6")
    return s
