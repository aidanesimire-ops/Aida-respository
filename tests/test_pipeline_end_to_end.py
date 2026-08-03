import json

import openpyxl
import pytest

from waterfront_engine.config import ConfigError, load_config
from waterfront_engine.demo import run_demo

CONFIG = "configs/fort_lauderdale.py"


@pytest.fixture(scope="module")
def demo(tmp_path_factory):
    out = tmp_path_factory.mktemp("demo")
    return run_demo(out, market_path=CONFIG), out


def test_demo_writes_both_workbooks_and_report(demo):
    paths, _ = demo
    assert paths["ownership_workbook"].exists()
    assert paths["condo_workbook"].exists()
    report = json.loads(paths["run_report"].read_text())
    assert report["market"] == "Fort Lauderdale, FL"
    assert report["classes"]["Single-Family"]["waterfront"] > 0
    assert report["beach_condo_units"] == 8


def test_ownership_workbook_tabs_and_headers(demo):
    paths, _ = demo
    book = openpyxl.load_workbook(paths["ownership_workbook"])
    for tab in ("Summary", "Single-Family", "Marina", "Medical", "Read Me"):
        assert tab in book.sheetnames

    sheet = book["Single-Family"]
    headers = [c.value for c in sheet[2]]
    for expected in (
        "Folio",
        "Situs Address",
        "Owner 1",
        "Mailing Address",
        "Absentee",
        "Entity",
        "Just/Market Value",
        "Waterfront Type",
        "Water Body",
        "Managers/Members",
    ):
        assert expected in headers
    assert sheet.freeze_panes == "A3"
    assert sheet.max_row > 2, "single-family tab should carry waterfront rows"


def test_summary_counts_match_tabs(demo):
    paths, _ = demo
    book = openpyxl.load_workbook(paths["ownership_workbook"])
    summary = book["Summary"]
    headers = [c.value for c in summary[2]]
    rows = {
        r[headers.index("Asset Class")]: r[headers.index("Parcels")]
        for r in summary.iter_rows(min_row=3, values_only=True)
        if r[0]
    }
    for tab, parcels in rows.items():
        assert book[tab].max_row - 2 == parcels, f"{tab} row count disagrees with Summary"


def test_condo_workbook_is_unit_level(demo):
    paths, _ = demo
    book = openpyxl.load_workbook(paths["condo_workbook"])
    assert book.sheetnames[:2] == ["By Building", "Units"]

    units = book["Units"]
    headers = [c.value for c in units[2]]
    assert {"Unit Address", "Building", "Unit #", "Owner 1", "Mailing Address"} <= set(headers)
    assert units.max_row - 2 == 8

    building = headers.index("Building")
    unit_no = headers.index("Unit #")
    values = list(units.iter_rows(min_row=3, values_only=True))
    assert {v[building] for v in values} == {"4300 N OCEAN BLVD"}
    assert all(v[unit_no] for v in values), "every unit row needs a unit number"


def test_enrichment_fills_entities_only(demo):
    paths, _ = demo
    book = openpyxl.load_workbook(paths["ownership_workbook"])
    sheet = book["Single-Family"]
    headers = [c.value for c in sheet[2]]
    entity = headers.index("Entity")
    managers = headers.index("Managers/Members")

    rows = list(sheet.iter_rows(min_row=3, values_only=True))
    assert rows, "expected waterfront single-family rows"
    for row in rows:
        if row[entity] == "Y":
            assert row[managers], "entity-owned rows should carry manager detail"
        else:
            assert not row[managers], "person-owned rows must not be enriched"


def test_qa_tabs_present(demo):
    paths, _ = demo
    book = openpyxl.load_workbook(paths["ownership_workbook"])
    assert "QA — Use Code Census" in book.sheetnames


def test_stage_files_written(demo):
    _, out = demo
    data = out / "data"
    assert (data / "parcels_all.gpkg").exists()
    assert (data / "water.gpkg").exists()
    assert (data / "class_single_family.gpkg").exists()
    assert (data / "waterfront_single_family.gpkg").exists()
    assert (data / "beach_condos.gpkg").exists()


# -- config loading ------------------------------------------------------
def test_load_config_shapes():
    cfg = load_config(CONFIG)
    assert cfg.market.parcel_query_url.endswith("/94/query")
    assert cfg.market.local_cities[0] == "FORT LAUDERDALE"
    assert "FT LAUDERDALE" in cfg.market.local_cities
    assert cfg.asset_classes["specialty_medical"].source_codes == ("73", "74", "17", "18", "19")
    assert "04" in cfg.all_source_codes  # condo codes are pulled too
    assert cfg.market.slug() == "fort_lauderdale_fl"


def test_missing_required_field_is_rejected(tmp_path):
    bad = tmp_path / "bad_market.py"
    bad.write_text(
        "MARKET = {'name': 'X', 'parcel_service': 'http://x', 'parcel_layer': 1,"
        " 'water_layer': 2, 'proj_crs': 2236, 'home_city': 'X', 'f_owner1': 'O'}\n"
        "ASSET_CLASSES = {'sf': {'label': 'SF', 'group': 'Residential', 'codes': ['01']}}\n"
    )
    with pytest.raises(ConfigError, match="required field"):
        load_config(bad)


def test_include_kw_from_without_pattern_is_rejected(tmp_path):
    bad = tmp_path / "bad_classes.py"
    bad.write_text(
        "from configs.fort_lauderdale import MARKET\n"
        "ASSET_CLASSES = {'med': {'label': 'M', 'group': 'S', 'codes': ['73'],"
        " 'include_kw_from': ['19']}}\n"
    )
    with pytest.raises(ConfigError, match="include_kw"):
        load_config(bad)
